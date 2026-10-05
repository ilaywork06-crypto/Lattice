"""Excel import/export, one sheet per template (built on openpyxl).

The import workbook has **one sheet per template**, and each sheet's header row
is exactly that template's item-creation fields — the per-item (grey) and list
fields, in the template's order — plus an optional *Serial* column (blank → the
next serial is issued) and, for containers, a *Contents* column listing the
serials of the items to place inside.

Importing is all-or-nothing: every row of every sheet is validated through the
same service that creates items from the UI; if any cell is wrong, nothing is
saved and the response lists **every** bad cell (sheet + A1 reference + reason).
"""

from __future__ import annotations

import io
import re

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func
from sqlalchemy.orm import Session

from lattice_core.errors import DomainError
from lattice_core.models import (
    CATALOG_FIELD_TYPES,
    CatalogOption,
    FieldMode,
    FieldType,
    Item,
    ItemState,
    ItemTemplate,
    ItemType,
    Location,
    TemplateField,
    User,
)
from lattice_core.services import items as item_svc
from lattice_core.services import serials as serials_svc
from lattice_core.services import views

META_SHEET = "_lattice"
README_SHEET = "Read me"
SERIAL_HEADER = "Serial"
CONTENTS_HEADER = "Contents"
_TYPE_ORDER = {ItemType.card: 0, ItemType.assembly: 1, ItemType.setup: 2}
_BAD_TITLE_CHARS = re.compile(r"[\[\]:*?/\\]")
_REQUIRED_FILL = PatternFill("solid", fgColor="FFE9E3F8")
_HEADER_FILL = PatternFill("solid", fgColor="FFEFF1F5")


def sheet_title(tpl: ItemTemplate) -> str:
    """Unique per template ((type, prefix) is unique) and within Excel's rules."""
    head = f"{serials_svc.format_serial(tpl.type, tpl.serial_prefix, 0)[:5]} "
    return (head + _BAD_TITLE_CHARS.sub(" ", tpl.name))[:31].strip()


def import_fields(tpl: ItemTemplate) -> list[TemplateField]:
    """The fields filled in when an item is created (files can't ride in a cell)."""
    return [
        f for f in tpl.fields
        if f.mode != FieldMode.fixed and f.field_type != FieldType.files
    ]


def _templates(db: Session, template_id: int | None, item_type: ItemType | None):
    q = db.query(ItemTemplate)
    if template_id is not None:
        q = q.filter(ItemTemplate.id == template_id)
    if item_type is not None:
        q = q.filter(ItemTemplate.type == item_type)
    return sorted(q.all(), key=lambda t: (_TYPE_ORDER[t.type], t.name.lower()))


def _allowed_values(db: Session, f: TemplateField) -> list[str]:
    options = (f.config or {}).get("options") or []
    if f.field_type == FieldType.enum or (f.mode == FieldMode.choice and options):
        return views.options_display(db, f)
    if f.field_type == FieldType.status:
        return [s.value for s in ItemState]
    if f.field_type == FieldType.letter:
        return [chr(c) for c in range(65, 91)]
    if f.field_type == FieldType.boolean:
        return ["yes", "no"]
    if f.field_type in CATALOG_FIELD_TYPES:
        return [
            o.value for o in db.query(CatalogOption).filter(
                CatalogOption.category == CATALOG_FIELD_TYPES[f.field_type],
                CatalogOption.active.is_(True),
            ).order_by(CatalogOption.sort_order, CatalogOption.value)
        ]
    if f.field_type == FieldType.location:
        return [loc.name for loc in db.query(Location).order_by(Location.name)]
    return []


def _hint(db: Session, f: TemplateField) -> str:
    parts = [f"Type: {f.field_type.value}", "Required" if f.required else "Optional"]
    pattern = (f.config or {}).get("pattern")
    if pattern:
        parts.append(f"Format {pattern}: type only the digits (#)")
    if f.field_type == FieldType.description:
        parts.append(f"At least {(f.config or {}).get('min_length', 8)} non-blank characters")
    if f.field_type == FieldType.managers:
        parts.append("Manager emails or names, separated by commas")
    if f.field_type == FieldType.responsible:
        parts.append("A user's email or name")
    if f.field_type == FieldType.parent:
        parts.append("The serial of the item this one goes into")
    if f.field_type == FieldType.date:
        parts.append("YYYY-MM-DD")
    if f.mode == FieldMode.choice:
        parts.append("Blank = the first value of the list")
    allowed = _allowed_values(db, f)
    if allowed:
        shown = ", ".join(allowed[:25]) + (" …" if len(allowed) > 25 else "")
        parts.append(f"Allowed: {shown}")
    return "\n".join(parts)


# ─────────────────────────── template workbook ───────────────────────────
def build_template(
    db: Session, template_id: int | None = None, item_type: ItemType | None = None
) -> bytes:
    templates = _templates(db, template_id, item_type)
    if not templates:
        raise DomainError("There are no templates to build an import file from")
    wb = Workbook()
    readme = wb.active
    readme.title = README_SHEET
    readme.append(["Lattice import"])
    readme["A1"].font = Font(bold=True, size=14)
    for line in [
        "One sheet per template. Each row creates one item from that template.",
        "The header row lists the fields filled in when an item is created; "
        "highlighted headers are required. Hover a header for its format.",
        f"'{SERIAL_HEADER}': leave blank to get the next serial automatically.",
        f"'{CONTENTS_HEADER}': serials of existing items (or items from earlier rows/"
        "sheets) to place inside, separated by commas.",
        "Fields set on the template itself are not listed — every item gets them.",
        "If any cell is invalid nothing is imported, and every invalid cell is listed.",
    ]:
        readme.append([line])
    readme.column_dimensions["A"].width = 110

    meta = wb.create_sheet(META_SHEET)
    meta.append(["sheet", "template_id", "type", "serial_prefix"])
    meta.sheet_state = "hidden"

    for tpl in templates:
        ws = wb.create_sheet(sheet_title(tpl))
        meta.append([ws.title, tpl.id, tpl.type.value, tpl.serial_prefix])
        headers: list[tuple[str, str]] = [
            (SERIAL_HEADER,
             f"Optional. Blank = next serial, e.g. {serials_svc.next_serial(db, tpl)}")
        ]
        fields = import_fields(tpl)
        headers += [(f.label, _hint(db, f)) for f in fields]
        if tpl.child_templates:
            names = ", ".join(c.name for c in tpl.child_templates)
            headers.append((CONTENTS_HEADER, f"Serials of items to place inside ({names})"))
        ws.append([h for h, _ in headers])
        for col, (_, hint) in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True)
            cell.fill = _HEADER_FILL
            cell.comment = Comment(hint, "Lattice")
            ws.column_dimensions[get_column_letter(col)].width = max(14, len(str(cell.value)) + 4)
        for col, f in enumerate(fields, start=2):
            if f.required:
                ws.cell(row=1, column=col).fill = _REQUIRED_FILL
            allowed = _allowed_values(db, f)
            formula = '"' + ",".join(a.replace('"', "") for a in allowed) + '"'
            if allowed and len(formula) <= 255 and f.field_type != FieldType.managers:
                dv = DataValidation(type="list", formula1=formula, allow_blank=not f.required)
                ws.add_data_validation(dv)
                letter = get_column_letter(col)
                dv.add(f"{letter}2:{letter}1000")
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─────────────────────────── import ───────────────────────────
def _norm_header(value) -> str:
    return re.sub(r"\s*\*$", "", str(value or "").strip()).lower()


def _resolve_sheets(db: Session, wb) -> list[tuple[object, ItemTemplate | None]]:
    mapping: dict[str, int] = {}
    if META_SHEET in wb.sheetnames:
        for row in wb[META_SHEET].iter_rows(min_row=2, values_only=True):
            if row and row[0] and row[1]:
                try:
                    mapping[str(row[0])] = int(row[1])
                except (TypeError, ValueError):
                    continue
    out = []
    for ws in wb.worksheets:
        if ws.title in (META_SHEET, README_SHEET):
            continue
        tpl = db.get(ItemTemplate, mapping[ws.title]) if ws.title in mapping else None
        if tpl is None:
            m = re.match(r"^([CAS])-([A-Za-z]{3})\b", ws.title)
            if m:
                item_type = {"C": ItemType.card, "A": ItemType.assembly, "S": ItemType.setup}[
                    m.group(1)
                ]
                tpl = (
                    db.query(ItemTemplate)
                    .filter(ItemTemplate.type == item_type,
                            ItemTemplate.serial_prefix == m.group(2).upper())
                    .first()
                )
        if tpl is None:
            tpl = db.query(ItemTemplate).filter(
                func.lower(ItemTemplate.name) == ws.title.strip().lower()
            ).first()
        out.append((ws, tpl))
    out.sort(key=lambda p: _TYPE_ORDER[p[1].type] if p[1] else -1)
    return out


def _is_blank_row(row) -> bool:
    return all(v is None or (isinstance(v, str) and not v.strip()) for v in row)


def import_items(db: Session, content: bytes, user: User) -> dict:  # noqa: C901
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse failure is "not a workbook"
        raise DomainError(f"This is not a readable Excel (.xlsx) file: {exc}") from exc

    errors: list[dict] = []
    created = 0
    by_template: dict[str, int] = {}
    savepoint = db.begin_nested()

    def err(sheet: str, row: int | None, col: int | None, message: str, column: str | None = None):
        letter = get_column_letter(col) if col else None
        errors.append({
            "sheet": sheet,
            "cell": f"{letter}{row}" if letter and row else None,
            "row": row,
            "column": column,
            "error": message,
        })

    sheets = _resolve_sheets(db, wb)
    if not sheets:
        raise DomainError("The workbook has no template sheets to import")

    for ws, tpl in sheets:
        if tpl is None:
            err(ws.title, None, None,
                "This sheet doesn't match any template — download a fresh import file")
            continue
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = rows[0]
        by_label = {f.label.strip().lower(): f for f in import_fields(tpl)}
        by_key = {f.key.lower(): f for f in import_fields(tpl)}
        columns: dict[int, TemplateField] = {}
        serial_col = contents_col = None
        for col, raw in enumerate(header, start=1):
            name = _norm_header(raw)
            if not name:
                continue
            if name == SERIAL_HEADER.lower():
                serial_col = col
            elif name == CONTENTS_HEADER.lower():
                contents_col = col
            elif name in by_label or name in by_key:
                columns[col] = by_label.get(name) or by_key[name]
            else:
                err(ws.title, 1, col, f"'{raw}' is not a field of '{tpl.name}'", str(raw))
        present = set(columns.values())
        for f in import_fields(tpl):
            if f.required and f not in present:
                err(ws.title, 1, None, f"The required column '{f.label}' is missing", f.label)
        if any(e["sheet"] == ws.title and e["row"] == 1 for e in errors):
            continue
        col_of_key = {f.key: c for c, f in columns.items()}

        for r, row in enumerate(rows[1:], start=2):
            if _is_blank_row(row):
                continue
            def cell(c, row=row):
                return row[c - 1] if c and c - 1 < len(row) else None

            values = {f.key: cell(c) for c, f in columns.items() if cell(c) not in (None, "")}

            serial = cell(serial_col)
            if serial not in (None, ""):
                try:
                    serials_svc.validate_manual(db, tpl, str(serial))
                except DomainError as exc:
                    err(ws.title, r, serial_col, str(exc), SERIAL_HEADER)
                    continue

            child_ids: list[int] = []
            contents = cell(contents_col)
            bad_contents = False
            if contents not in (None, ""):
                for s in re.split(r"[,;\n]", str(contents)):
                    s = s.strip().upper()
                    if not s:
                        continue
                    child = db.query(Item).filter(Item.serial == s).first()
                    if child is None:
                        err(ws.title, r, contents_col, f"no item has the serial '{s}'",
                            CONTENTS_HEADER)
                        bad_contents = True
                    else:
                        child_ids.append(child.id)

            row_sp = db.begin_nested()
            try:
                item = item_svc.create_item(
                    db,
                    {"template_id": tpl.id, "values": values,
                     "serial": str(serial) if serial not in (None, "") else None},
                    user,
                )
                db.flush()
            except DomainError as exc:
                row_sp.rollback()
                if exc.errors:
                    for e in exc.errors:
                        c = col_of_key.get(e.get("field"))
                        err(ws.title, r, c, e["error"], e.get("label") or e.get("field"))
                else:
                    err(ws.title, r, None, str(exc))
                continue
            for cid in child_ids:
                try:
                    item_svc.link_item(db, db.get(Item, cid), item, user)
                except DomainError as exc:
                    err(ws.title, r, contents_col, str(exc), CONTENTS_HEADER)
                    bad_contents = True
            if bad_contents:
                row_sp.rollback()
                continue
            row_sp.commit()
            created += 1
            by_template[tpl.name] = by_template.get(tpl.name, 0) + 1

    if errors:
        savepoint.rollback()
        db.rollback()
        raise DomainError(
            f"Nothing was imported: {len(errors)} problem(s) found. Fix the listed cells "
            "and import the file again.",
            errors,
        )
    savepoint.commit()
    db.commit()
    return {"created": created, "by_template": by_template, "errors": []}


# ─────────────────────────── export ───────────────────────────
def _cell_value(display):
    if display is None:
        return None
    if isinstance(display, list):
        return ", ".join(
            str(d.get("name") if isinstance(d, dict) else d) for d in display
        )
    if isinstance(display, bool):
        return "yes" if display else "no"
    return display


def export_items(
    db: Session, template_id: int | None = None, item_type: ItemType | None = None
) -> bytes:
    """Every item, one sheet per template, readable names instead of ids."""
    templates = _templates(db, template_id, item_type)
    wb = Workbook()
    wb.remove(wb.active)
    for tpl in templates:
        ws = wb.create_sheet(sheet_title(tpl))
        headers = [SERIAL_HEADER, "State", "Location", "Parent"]
        headers += [f.label for f in tpl.fields if f.field_type not in (
            FieldType.status, FieldType.location, FieldType.parent
        )]
        if tpl.child_templates:
            headers.append(CONTENTS_HEADER)
        headers += ["Created", "Updated"]
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = _HEADER_FILL
        items = (
            db.query(Item).filter(Item.template_id == tpl.id).order_by(Item.serial).all()
        )
        for it in items:
            row = [
                it.serial,
                it.state.value,
                it.location.name if it.location else None,
                it.parent.serial if it.parent else None,
            ]
            for f in tpl.fields:
                if f.field_type in (FieldType.status, FieldType.location, FieldType.parent):
                    continue
                value = item_svc.effective_value(db, it, f)
                row.append(_cell_value(views.display_value(db, f.field_type, value)))
            if tpl.child_templates:
                row.append(", ".join(c.serial for c in it.children) or None)
            row += [
                it.created_at.replace(tzinfo=None) if it.created_at else None,
                it.updated_at.replace(tzinfo=None) if it.updated_at else None,
            ]
            ws.append(row)
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 18
        ws.freeze_panes = "A2"
    if not wb.worksheets:
        ws = wb.create_sheet("items")
        ws.append(["No templates yet"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
