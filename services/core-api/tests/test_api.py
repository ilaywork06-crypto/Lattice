"""API tests for the template-driven model.

The session-scoped world (see conftest.py) is shared, so tests that change it
create their own items rather than editing the fixtures' ones.
"""

import io

from openpyxl import Workbook, load_workbook


def _tpl(client, admin, prefix):
    for t in client.get("/templates", headers=admin).json():
        if t["serial_prefix"] == prefix:
            return t
    raise AssertionError(prefix)


def _new(client, headers, template_id, expect=201, **values):
    r = client.post("/items", json={"template_id": template_id, "values": values},
                    headers=headers)
    assert r.status_code == expect, r.text
    return r.json()


def _loc(client, admin, name):
    return next(loc for loc in client.get("/locations", headers=admin).json()
                if loc["name"] == name)


def _users(client, admin):
    return {u["email"]: u for u in client.get("/users", headers=admin).json()}


# ─────────────────────────── basics & roles ───────────────────────────
def test_health(client):
    assert client.get("/health").json()["status"] == "ok"


def test_auth_and_roles(client, viewer):
    me = client.get("/auth/me", headers=viewer).json()
    assert me["role"] == "viewer"


def test_viewer_and_editor_cannot_mutate_directly(client, viewer, editor, a_setup):
    for headers in (viewer, editor):
        assert client.patch(f"/items/{a_setup}", json={"values": {}},
                            headers=headers).status_code == 403
        assert client.post("/templates", json={}, headers=headers).status_code in (403, 422)


# ─────────────────────────── templates ───────────────────────────
def _template_body(**over):
    body = {"type": "card", "name": "Test Card", "card_type": "white",
            "serial_prefix": "TST", "fields": []}
    body.update(over)
    return body


def test_template_rules_are_enforced(client, admin):
    cases = [
        (_template_body(serial_prefix="T1X"), "three Latin letters"),
        (_template_body(serial_prefix="PRB"), "already used"),
        (_template_body(name="power regulator board"), "already exists"),
        (_template_body(card_type=None), "card type"),
        (_template_body(type="setup"), "Only card templates"),
        (_template_body(fields=[{"label": "Q", "field_type": "quantity"}]), "commercial"),
        (_template_body(fields=[{"label": "Where", "field_type": "location",
                                 "mode": "fixed"}]), "per item"),
        (_template_body(fields=[{"label": "Kind", "field_type": "enum"}]), "at least one"),
        (_template_body(fields=[{"label": "Pick", "field_type": "text", "mode": "choice"}]),
         "list field needs"),
        (_template_body(fields=[{"label": "Owner", "field_type": "text", "mode": "fixed",
                                 "required": True}]), "needs its value"),
        (_template_body(fields=[{"label": "A", "field_type": "team"},
                                {"label": "B", "field_type": "team"}]), "only one team"),
        (_template_body(fields=[{"label": "A", "field_type": "text"},
                                {"label": "a", "field_type": "text"}]), "two fields"),
        (_template_body(fields=[{"label": "Code", "field_type": "string",
                                 "config": {"pattern": "XX-"}}]), "#"),
    ]
    for body, message in cases:
        r = client.post("/templates", json=body, headers=admin)
        assert r.status_code == 400, (body, r.text)
        assert message.lower() in r.json()["detail"].lower(), (message, r.text)


def test_template_children_respect_the_hierarchy(client, admin, templates):
    r = client.post("/templates", json={
        "type": "assembly", "name": "Bad Assembly", "serial_prefix": "BAD",
        "child_template_ids": [templates["FTS"]],
    }, headers=admin)
    assert r.status_code == 400
    assert "cannot contain setup" in r.json()["detail"]


def test_template_detail_exposes_fields_children_and_next_serial(client, admin, templates):
    t = client.get(f"/templates/{templates['SPM']}", headers=admin).json()
    assert [c["serial_prefix"] for c in t["child_templates"]] == ["PRB", "FPG"]
    assert [p["serial_prefix"] for p in t["parent_templates"]] == ["FTS"]
    team = next(f for f in t["fields"] if f["key"] == "team")
    assert team["mode"] == "fixed" and team["fixed_display"] == "HW-Team-A"
    assert t["next_serial"] == "A-SPM-002"


def test_grouped_counts_exclude_destroyed(client, admin, templates):
    tid = templates["FPG"]
    before = next(t for t in client.get("/templates?type=card", headers=admin).json()
                  if t["id"] == tid)["counts"]
    item = _new(client, admin, tid, description="Spare board for testing")
    client.post(f"/items/{item['id']}/state", json={"state": "destroyed"}, headers=admin)
    after = next(t for t in client.get("/templates?type=card", headers=admin).json()
                 if t["id"] == tid)["counts"]
    assert after["total"] == before["total"]
    assert after["destroyed"] == before["destroyed"] + 1


def test_items_of_a_template_list_state_location_parent_serial(client, admin, templates):
    rows = client.get(f"/items?template_id={templates['PRB']}", headers=admin).json()
    assert all(r["template_id"] == templates["PRB"] for r in rows)
    assembled = next(r for r in rows if r["serial"] == "C-PRB-001")
    assert assembled["parent_label"] == "Signal Processing Module (A-SPM-001)"
    assert assembled["location_name"] == "Integration Hall"
    assert assembled["storage_status"] == "assembled"


# ─────────────────────────── creation from a template ───────────────────────────
def test_items_are_created_only_from_a_template(client, admin):
    assert client.post("/items", json={"values": {}}, headers=admin).status_code == 422
    r = client.post("/items", json={"template_id": 9999}, headers=admin)
    assert r.status_code == 400 and "template" in r.json()["detail"]


def test_serials_are_issued_per_template_and_unique(client, admin, templates):
    tid = templates["NIC"]
    a = _new(client, admin, tid, quantity=2)
    b = _new(client, admin, tid, quantity=1)
    num = lambda s: int(s.rsplit("-", 1)[1])  # noqa: E731
    assert a["serial"].startswith("C-NIC-") and num(b["serial"]) == num(a["serial"]) + 1

    manual = client.post("/items", json={"template_id": tid, "serial": "c-nic-050",
                                         "values": {"quantity": 1}}, headers=admin)
    assert manual.status_code == 201 and manual.json()["serial"] == "C-NIC-050"
    nxt = _new(client, admin, tid, quantity=1)
    assert nxt["serial"] == "C-NIC-051", "numbering continues above the highest"

    for bad, why in (("C-NIC-050", "already used"), ("A-NIC-099", "starts with 'C-'"),
                     ("C-PRB-099", "prefix 'NIC'"), ("NIC-1", "not a valid serial")):
        r = client.post("/items", json={"template_id": tid, "serial": bad,
                                        "values": {"quantity": 1}}, headers=admin)
        assert r.status_code == 400 and why in r.json()["detail"], (bad, r.text)

    # a manual edit keeps the scheme and refuses duplicates too
    r = client.patch(f"/items/{a['id']}", json={"serial": "C-NIC-050"}, headers=admin)
    assert r.status_code == 400
    r = client.patch(f"/items/{a['id']}", json={"serial": "C-NIC-777"}, headers=admin)
    assert r.status_code == 200 and r.json()["serial"] == "C-NIC-777"


def test_new_items_are_built_unless_a_status_is_chosen(client, admin, templates):
    tid = templates["PRB"]
    default = _new(client, admin, tid)
    assert default["state"] == "built"
    chosen = _new(client, admin, tid, status="ok")
    assert chosen["state"] == "ok"
    r = client.post("/items", json={"template_id": tid, "values": {"status": "faulty"}},
                    headers=admin)
    assert r.status_code == 400, "faulty is not in this template's list"


def test_list_field_defaults_to_the_head_and_refuses_other_values(client, admin, templates):
    tid = templates["PRB"]
    item = _new(client, admin, tid)
    version = next(f for f in item["fields"] if f["key"] == "version")
    assert version["value"] == "1.3", "the first list value is the default"
    other = _new(client, admin, tid, version="1.2")
    assert next(f for f in other["fields"] if f["key"] == "version")["value"] == "1.2"
    r = client.post("/items", json={"template_id": tid, "values": {"version": "9.9"}},
                    headers=admin)
    assert r.status_code == 400


def test_template_values_cannot_be_overridden_per_item(client, admin, templates):
    r = client.post("/items", json={"template_id": templates["PRB"],
                                    "values": {"project": "Sparrow"}}, headers=admin)
    assert r.status_code == 400
    errors = r.json()["errors"]
    assert errors[0]["field"] == "project" and "template" in errors[0]["error"]


def test_every_invalid_field_is_reported_at_once(client, admin, templates):
    r = client.post("/items", json={
        "template_id": templates["FPG"],
        "values": {"description": "short", "board_id": "12", "letter": "ab", "nope": 1},
    }, headers=admin)
    assert r.status_code == 400
    fields = {e["field"] for e in r.json()["errors"]}
    assert {"description", "board_id", "letter", "nope"} <= fields


def test_formats_letters_and_descriptions(client, admin, templates):
    item = _new(client, admin, templates["FPG"], description="   eight chars!  ",
                board_id="54321", letter="q")
    values = {f["key"]: f["value"] for f in item["fields"]}
    assert values["board_id"] == "FP-54321", "the fixed letters are filled in automatically"
    assert values["letter"] == "Q"
    assert values["description"] == "eight chars!"
    # the full form is accepted as well
    again = _new(client, admin, templates["FPG"], description="another board",
                 board_id="FP-11111")
    assert {f["key"]: f["value"] for f in again["fields"]}["board_id"] == "FP-11111"


def test_new_cards_default_to_the_desiccator(client, admin, templates):
    item = _new(client, admin, templates["NIC"], quantity=1)
    assert item["location"]["is_desiccator"] is True
    assert item["storage_status"] == "desiccator"


def test_only_commercial_cards_carry_a_quantity(client, admin, templates):
    r = client.post("/items", json={"template_id": templates["FPG"],
                                    "values": {"description": "board", "quantity": 3}},
                    headers=admin)
    assert r.status_code == 400  # not a field of a serialised card's template


def test_managers_field_accepts_managers_only(client, admin, templates):
    users = _users(client, admin)
    r = client.post("/templates", json={
        "type": "setup", "name": "Mgr Setup", "serial_prefix": "MGR",
        "fields": [{"label": "Managers", "field_type": "managers", "mode": "fixed",
                    "fixed_value": [users["dana@lattice.io"]["id"]]}],
    }, headers=admin)
    assert r.status_code == 400 and "not a manager" in r.json()["detail"]


def test_fixed_managers_link_every_item_and_follow_template_edits(client, admin, templates):
    users = _users(client, admin)
    tid = templates["PRB"]
    tpl = client.get(f"/templates/{tid}", headers=admin).json()
    items = client.get(f"/items?template_id={tid}", headers=admin).json()
    assert all(i["manager_names"] == ["Noa (Team Lead)"] for i in items)

    fields = [{**f} for f in tpl["fields"]]
    for f in fields:
        if f["key"] == "managers":
            f["fixed_value"] = [users["admin@lattice.io"]["id"]]
    r = client.patch(f"/templates/{tid}", json={"fields": fields}, headers=admin)
    assert r.status_code == 200, r.text
    items = client.get(f"/items?template_id={tid}", headers=admin).json()
    assert all(i["manager_names"] == ["System Administrator"] for i in items), \
        "editing the template moved every item under the new manager"

    # restore for the other tests
    for f in fields:
        if f["key"] == "managers":
            f["fixed_value"] = [users["noa@lattice.io"]["id"]]
    assert client.patch(f"/templates/{tid}", json={"fields": fields},
                        headers=admin).status_code == 200


def test_switching_a_fixed_field_to_per_item_keeps_each_items_value(client, admin):
    r = client.post("/templates", json={
        "type": "setup", "name": "Mode Switch", "serial_prefix": "MSW",
        "fields": [{"label": "Owner", "field_type": "text", "mode": "fixed",
                    "fixed_value": "Lab"}],
    }, headers=admin)
    tpl = r.json()
    item = _new(client, admin, tpl["id"])
    assert item["fields"][0]["value"] == "Lab"

    field = {**tpl["fields"][0], "fixed_value": "Lab 2"}
    client.patch(f"/templates/{tpl['id']}", json={"fields": [field]}, headers=admin)
    assert client.get(f"/items/{item['id']}", headers=admin).json()["fields"][0]["value"] \
        == "Lab 2", "a template value is read through, so editing it changes every item"

    field = {**field, "mode": "item"}
    client.patch(f"/templates/{tpl['id']}", json={"fields": [field]}, headers=admin)
    got = client.get(f"/items/{item['id']}", headers=admin).json()["fields"][0]
    assert got["mode"] == "item" and got["value"] == "Lab 2"


def test_a_field_type_cannot_change_in_place(client, admin, templates):
    tpl = client.get(f"/templates/{templates['FTS']}", headers=admin).json()
    fields = [{**f} for f in tpl["fields"]]
    fields[-1]["field_type"] = "integer"
    r = client.patch(f"/templates/{templates['FTS']}", json={"fields": fields}, headers=admin)
    assert r.status_code == 400 and "cannot change" in r.json()["detail"]


def test_a_template_with_items_cannot_be_deleted(client, admin, templates):
    r = client.delete(f"/templates/{templates['PRB']}", headers=admin)
    assert r.status_code == 400
    empty = client.post("/templates", json={"type": "setup", "name": "Empty",
                                            "serial_prefix": "EMP"}, headers=admin).json()
    assert client.delete(f"/templates/{empty['id']}", headers=admin).status_code == 204


def test_editors_propose_template_edits_and_managers_apply_them(
    client, admin, editor, templates
):
    tid = templates["SPM"]
    assert client.patch(f"/templates/{tid}", json={"description": "x"},
                        headers=editor).status_code == 403
    r = client.post("/change-requests", json={
        "action": "template_update", "template_id": tid,
        "payload": {"description": "Signal processing, rev B"},
        "reason": "Clarify what the module is",
    }, headers=editor)
    assert r.status_code == 201, r.text
    cr = r.json()
    assert cr["template_id"] == tid and cr["description"]
    assert client.post(f"/change-requests/{cr['id']}/approve", headers=admin).status_code == 200
    assert client.get(f"/templates/{tid}", headers=admin).json()["description"] == \
        "Signal processing, rev B"


# ─────────────────────────── hierarchy ───────────────────────────
def test_links_follow_the_templates(client, admin, templates):
    nic = _new(client, admin, templates["NIC"], quantity=1)
    spm = _new(client, admin, templates["SPM"])
    r = client.post(f"/items/{nic['id']}/link", json={"parent_id": spm["id"]}, headers=admin)
    assert r.status_code == 400 and "don't include" in r.json()["detail"]
    setup = _new(client, admin, templates["FTS"],
                 location=_loc(client, admin, "Lab B — Bench 4")["id"])
    r = client.post(f"/items/{nic['id']}/link", json={"parent_id": setup["id"]}, headers=admin)
    assert r.status_code == 200
    assert r.json()["location"]["name"] == "Lab B — Bench 4"
    assert r.json()["storage_status"] == "assembled"


def test_move_cascades_and_linked_items_cannot_move_alone(client, admin, templates):
    lab_a = _loc(client, admin, "Lab A — Bench 1")["id"]
    storage = _loc(client, admin, "Storage Room")["id"]
    spm = _new(client, admin, templates["SPM"], location=lab_a)
    card = _new(client, admin, templates["FPG"], description="cascade test card")
    client.put(f"/items/{spm['id']}/children", json={"child_ids": [card["id"]]}, headers=admin)
    r = client.post(f"/items/{spm['id']}/move", json={"location_id": storage}, headers=admin)
    assert r.status_code == 200
    moved = client.get(f"/items/{card['id']}", headers=admin).json()
    assert moved["location_id"] == storage
    r = client.post(f"/items/{card['id']}/move", json={"location_id": lab_a}, headers=admin)
    assert r.status_code == 400 and "sits inside" in r.json()["detail"]

    # unlinking can say where the card now is
    desiccator = _loc(client, admin, "Desiccator — Team B")["id"]
    r = client.post(f"/items/{card['id']}/unlink", json={"location_id": desiccator},
                    headers=admin)
    assert r.json()["storage_status"] == "desiccator"


def test_create_with_parent_and_children(client, admin, templates):
    setup = _new(client, admin, templates["FTS"],
                 location=_loc(client, admin, "Lab A — Bench 1")["id"])
    card = _new(client, admin, templates["FPG"], description="child card one")
    r = client.post("/items", json={"template_id": templates["SPM"], "values": {},
                                    "child_ids": [card["id"]]}, headers=admin)
    spm = r.json()
    assert [c["id"] for c in spm["children"]] == [card["id"]]
    r = client.post(f"/items/{spm['id']}/link", json={"parent_id": setup["id"]}, headers=admin)
    assert r.status_code == 200
    assert client.get(f"/items/{card['id']}", headers=admin).json()["location"]["name"] == \
        "Lab A — Bench 1", "linking a container drags its contents along"


def test_state_machine_requires_note_for_fault(client, admin, templates):
    item = _new(client, admin, templates["FTS"],
                location=_loc(client, admin, "Lab A — Bench 1")["id"])
    r = client.post(f"/items/{item['id']}/state", json={"state": "faulty"}, headers=admin)
    assert r.status_code == 400
    r = client.post(f"/items/{item['id']}/state", json={"state": "faulty", "note": "smoke"},
                    headers=admin)
    assert r.status_code == 200
    history = r.json()["state_history"]
    assert history[-1]["changed_by_name"] == "System Administrator"
    assert history[-1]["changed_at"]


def test_physical_state_is_not_edited_through_the_edit_form(client, admin, templates):
    item = _new(client, admin, templates["PRB"])
    r = client.patch(f"/items/{item['id']}", json={"values": {"location": 1}}, headers=admin)
    assert r.status_code == 400 and "move" in r.json()["detail"]


def test_a_no_op_save_writes_no_audit_entry(client, admin, templates):
    item = _new(client, admin, templates["SPM"], rev="C")
    before = len(client.get(f"/audit?item_id={item['id']}", headers=admin).json())
    client.patch(f"/items/{item['id']}", json={"values": {"rev": "C"}}, headers=admin)
    assert len(client.get(f"/audit?item_id={item['id']}", headers=admin).json()) == before
    client.patch(f"/items/{item['id']}", json={"values": {"rev": "D"}}, headers=admin)
    assert len(client.get(f"/audit?item_id={item['id']}", headers=admin).json()) == before + 1


def test_bulk_move_is_atomic(client, admin, templates):
    loose = _new(client, admin, templates["SPM"])
    linked = client.get("/items?type=card", headers=admin).json()
    linked = next(i for i in linked if i["parent_id"])
    storage = _loc(client, admin, "Storage Room")["id"]
    r = client.post("/items/bulk", json={"action": "move", "item_ids": [loose["id"],
                                         linked["id"]], "location_id": storage}, headers=admin)
    assert r.status_code == 400
    assert client.get(f"/items/{loose['id']}", headers=admin).json()["location_id"] != storage


# ─────────────────────────── change requests ───────────────────────────
def test_viewer_may_propose_a_location_change_only(client, viewer, admin, templates):
    item = _new(client, admin, templates["SPM"])
    storage = _loc(client, admin, "Storage Room")["id"]
    r = client.post("/change-requests", json={
        "action": "move", "item_id": item["id"], "payload": {"location_id": storage},
        "reason": "It was carried to storage",
    }, headers=viewer)
    assert r.status_code == 201, r.text
    cr = r.json()
    assert "Storage Room" in cr["description"], "the 'what' is generated, only 'why' is asked"
    r = client.post("/change-requests", json={
        "action": "state_change", "item_id": item["id"], "payload": {"state": "ok"},
        "reason": "x",
    }, headers=viewer)
    assert r.status_code == 403
    assert client.post(f"/change-requests/{cr['id']}/approve",
                       headers=viewer).status_code == 403
    assert client.post(f"/change-requests/{cr['id']}/approve", headers=admin).status_code == 200
    assert client.get(f"/items/{item['id']}", headers=admin).json()["location_id"] == storage


def test_editor_proposes_an_item_and_approval_creates_it(client, editor, admin, templates):
    before = len(client.get(f"/items?template_id={templates['SPM']}", headers=admin).json())
    r = client.post("/change-requests", json={
        "action": "create", "payload": {"template_id": templates["SPM"],
                                        "values": {"rev": "E"}},
        "reason": "Building a second module",
    }, headers=editor)
    assert r.status_code == 201, r.text
    assert r.json()["item_name"] == "Signal Processing Module"
    assert client.post(f"/change-requests/{r.json()['id']}/approve",
                       headers=admin).status_code == 200
    after = len(client.get(f"/items?template_id={templates['SPM']}", headers=admin).json())
    assert after == before + 1


def test_malformed_proposal_is_rejected_not_a_500(client, editor, admin, a_setup):
    r = client.post("/change-requests", json={
        "action": "state_change", "item_id": a_setup, "payload": {"state": "bogus"},
        "reason": "x",
    }, headers=editor)
    r = client.post(f"/change-requests/{r.json()['id']}/approve", headers=admin)
    assert r.status_code == 400 and "not valid" in r.json()["detail"]


def test_pending_fixture_request_lists(client, admin):
    pending = client.get("/change-requests?status=pending", headers=admin).json()
    assert any(cr["reason"] == "It passed the bench test." for cr in pending)


# ─────────────────────────── catalog ───────────────────────────
def test_team_is_a_catalog_category(client, admin):
    teams = client.get("/catalog?category=team", headers=admin).json()
    assert {t["value"] for t in teams} >= {"HW-Team-A", "Integration"}
    hw = next(t for t in teams if t["value"] == "HW-Team-A")
    assert hw["usage_count"] >= 1, "assemblies made from SPM carry the team"


def test_catalog_links_are_two_way_and_cross_category(client, admin):
    opts = {o["value"]: o for o in client.get("/catalog", headers=admin).json()}
    team, space, defense = opts["Integration"], opts["Space"], opts["Defense"]
    r = client.put(f"/catalog/{team['id']}/links",
                   json={"category": "industry", "option_ids": [space["id"], defense["id"]]},
                   headers=admin)
    assert r.status_code == 200
    assert set(r.json()["linked_ids"]) >= {space["id"], defense["id"]}
    # seen from the industry's side, without touching it
    space_now = next(o for o in client.get("/catalog", headers=admin).json()
                     if o["id"] == space["id"])
    assert team["id"] in space_now["linked_ids"]
    # editing one category leaves the others alone
    assert opts["Falcon"]["id"] in r.json()["linked_ids"]
    # removing from the other end
    client.put(f"/catalog/{space['id']}/links", json={"category": "team", "option_ids": []},
               headers=admin)
    team_now = next(o for o in client.get("/catalog", headers=admin).json()
                    if o["id"] == team["id"])
    assert space["id"] not in team_now["linked_ids"]
    r = client.put(f"/catalog/{team['id']}/links", json={"category": "team",
                   "option_ids": []}, headers=admin)
    assert r.status_code == 400


def test_catalog_values_in_use_are_protected_and_renames_follow(client, admin, templates):
    opts = {o["value"]: o for o in client.get("/catalog", headers=admin).json()}
    assert client.delete(f"/catalog/{opts['Falcon']['id']}", headers=admin).status_code == 400
    r = client.patch(f"/catalog/{opts['HW-Team-A']['id']}", json={"value": "HW Team A"},
                     headers=admin)
    assert r.status_code == 200
    spm = client.get(f"/items?template_id={templates['SPM']}", headers=admin).json()[0]
    assert spm["team"] == "HW Team A", "items hold the id, so a rename reaches them"
    client.patch(f"/catalog/{opts['HW-Team-A']['id']}", json={"value": "HW-Team-A"},
                 headers=admin)


def test_catalog_management_is_manager_only(client, editor):
    r = client.post("/catalog", json={"category": "team", "value": "X"}, headers=editor)
    assert r.status_code == 403


# ─────────────────────────── desiccator & inventory ───────────────────────────
def test_the_desiccator_is_a_set_of_locations(client, admin, editor):
    locs = {loc["name"]: loc for loc in client.get("/locations", headers=admin).json()}
    assert {n for n, loc in locs.items() if loc["is_desiccator"]} == {
        "Desiccator — Team A", "Desiccator — Team B"}
    r = client.put("/locations/desiccator", json={"location_ids": [
        locs["Desiccator — Team A"]["id"], locs["Desiccator — Team B"]["id"],
        locs["Storage Room"]["id"]]}, headers=editor)
    assert r.status_code == 403
    r = client.patch(f"/locations/{locs['Storage Room']['id']}",
                     json={"is_desiccator": True}, headers=editor)
    assert r.status_code == 403


def test_available_stock_is_built_or_ok_loose_in_the_desiccator(client, admin, templates):
    def group():
        return next(g for g in client.get("/inventory/cards", headers=admin).json()
                    if g["template_id"] == tid)

    r = client.post("/templates", json={"type": "card", "name": "Stock Probe",
                                        "card_type": "factory", "serial_prefix": "STK"},
                    headers=admin)
    tid = r.json()["id"]
    a = _new(client, admin, tid)
    b = _new(client, admin, tid)
    g = group()
    assert (g["available"], g["desiccator"], g["total"]) == (2, 2, 2)

    client.post(f"/items/{a['id']}/state", json={"state": "faulty", "note": "dead"},
                headers=admin)
    g = group()
    assert g["available"] == 1 and g["desiccator"] == 2 and g["faulty"] == 1

    storage = _loc(client, admin, "Storage Room")["id"]
    client.post(f"/items/{b['id']}/move", json={"location_id": storage}, headers=admin)
    g = group()
    assert g["available"] == 0 and g["in_use"] == 1, "outside the desiccator = in use"

    # making the storage room part of the desiccator brings it back
    desiccator_ids = [loc["id"] for loc in client.get("/locations", headers=admin).json()
                      if loc["is_desiccator"]]
    client.put("/locations/desiccator", json={"location_ids": desiccator_ids + [storage]},
               headers=admin)
    assert group()["available"] == 1
    client.put("/locations/desiccator", json={"location_ids": desiccator_ids}, headers=admin)

    client.post(f"/items/{a['id']}/state", json={"state": "destroyed", "note": "scrapped"},
                headers=admin)
    assert group()["total"] == 1, "destroyed units are not inventory"


def test_thresholds_watch_a_template_and_alert(client, admin, templates, monkeypatch):
    published = []

    async def fake_publish(event):
        published.append(event)

    monkeypatch.setattr("lattice_core.events.publish_event", fake_publish)
    r = client.post("/inventory/thresholds", json={"template_id": templates["PRB"],
                                                   "min_quantity": 50}, headers=admin)
    assert r.status_code == 201 and r.json()["min_quantity"] == 50
    assert r.json()["is_low"] and r.json()["current_quantity"] < 50
    low = client.get("/inventory/low-stock", headers=admin).json()
    assert any(t["template_id"] == templates["PRB"] for t in low)
    events = [e for e in published if e.type.value == "inventory.low_stock"]
    assert events
    comp = next(c for e in events for c in e.payload["components"]
                if c["template_id"] == templates["PRB"])
    assert comp["min_quantity"] == 50 and comp["link"].startswith("/cards?template=")
    client.post("/inventory/thresholds", json={"template_id": templates["PRB"],
                                               "min_quantity": 2}, headers=admin)
    r = client.post("/inventory/thresholds", json={"template_id": templates["SPM"]},
                    headers=admin)
    assert r.status_code == 400


def test_dashboard_tiles_match_the_inventory(client, admin):
    s = client.get("/inventory/summary", headers=admin).json()
    groups = client.get("/inventory/cards", headers=admin).json()
    assert s["cards"] == sum(g["total"] for g in groups)
    assert s["cards_available"] == sum(g["available"] for g in groups)


def test_a_location_with_items_cannot_be_deleted(client, admin):
    integ = _loc(client, admin, "Integration Hall")
    assert client.delete(f"/locations/{integ['id']}", headers=admin).status_code == 400


def test_duplicate_location_names_are_refused(client, admin):
    r = client.post("/locations", json={"name": " lab a — bench 1 "}, headers=admin)
    assert r.status_code == 400


# ─────────────────────────── graphs ───────────────────────────
def test_template_graph(client, admin, templates):
    g = client.get("/graph/templates", headers=admin).json()
    edges = {(e["source"], e["target"]) for e in g["edges"]}
    assert (templates["FTS"], templates["SPM"]) in edges
    assert (templates["SPM"], templates["PRB"]) in edges
    sub = client.get(f"/graph/templates?root_template_id={templates['SPM']}",
                     headers=admin).json()
    assert {n["id"] for n in sub["nodes"]} == {templates["SPM"], templates["PRB"],
                                              templates["FPG"]}


def test_graphs_of_a_template_and_of_one_item(client, admin, templates):
    g = client.get(f"/graph?template_id={templates['FTS']}", headers=admin).json()
    assert g["roots"] and all(
        n["template_id"] == templates["FTS"] for n in g["nodes"] if n["id"] in g["roots"]
    )
    card = next(i for i in client.get(f"/items?template_id={templates['FPG']}",
                                      headers=admin).json() if i["serial"] == "C-FPG-001")
    g = client.get(f"/graph?root_id={card['id']}&ancestors=true", headers=admin).json()
    serials = {n["serial"] for n in g["nodes"]}
    assert {"C-FPG-001", "A-SPM-001", "S-FTS-001"} <= serials
    assert len(g["edges"]) >= 2


# ─────────────────────────── audit ───────────────────────────
def test_audit_periods_and_my_items(client, admin, manager, editor, templates):
    for period in ("day", "week", "month", "half_year", "year", "all"):
        assert client.get(f"/audit?period={period}", headers=admin).status_code == 200
    assert client.get("/audit?period=decade", headers=admin).status_code == 422

    mine = client.get("/audit/my-items?period=all", headers=manager).json()
    assert mine
    prb_ids = {i["id"] for i in client.get(f"/items?template_id={templates['PRB']}",
                                           headers=admin).json()}
    spm_ids = {i["id"] for i in client.get(f"/items?template_id={templates['SPM']}",
                                           headers=admin).json()}
    ids = {a["item_id"] for a in mine}
    assert ids & prb_ids and not ids & spm_ids, "only items Noa is linked to"
    assert all(a["item_id"] for a in mine), "no account/template administration noise"
    # someone linked to nothing sees nothing — not everything
    assert client.get("/audit/my-items?period=all", headers=editor).json() == []


def test_audit_export_is_an_excel_file(client, admin, a_setup):
    r = client.get(f"/audit/export?item_id={a_setup}&period=all", headers=admin)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "When (UTC)"
    assert ws.max_row > 2


# ─────────────────────────── documents ───────────────────────────
def test_documents_are_real_uploads(client, admin, viewer, a_setup):
    r = client.post(f"/items/{a_setup}/documents",
                    files={"file": ("report.pdf", b"%PDF-1.4 hello", "application/pdf")},
                    data={"doc_type": "test"}, headers=admin)
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["is_file"] and doc["size_bytes"] == 14 and doc["name"] == "report.pdf"
    dl = client.get(f"/documents/{doc['id']}/download", headers=viewer)
    assert dl.status_code == 200 and dl.content == b"%PDF-1.4 hello"
    assert client.get(f"/documents/{doc['id']}/download").status_code == 401
    link = client.post(f"/items/{a_setup}/documents",
                       data={"name": "Wiki", "url": "https://wiki.local/x"}, headers=admin)
    assert link.status_code == 201 and not link.json()["is_file"]
    assert client.delete(f"/items/{a_setup}/documents/{doc['id']}",
                         headers=admin).status_code == 204
    assert client.get(f"/documents/{doc['id']}/download", headers=admin).status_code == 404


def test_a_files_field_takes_staged_uploads(client, admin, editor, templates):
    staged = client.post("/uploads", files={"file": ("ds.txt", b"datasheet", "text/plain")},
                         headers=editor)
    assert staged.status_code == 201
    item = _new(client, admin, templates["FPG"], description="board with sheet",
                datasheet=[staged.json()["id"]])
    files = next(f for f in item["fields"] if f["key"] == "datasheet")
    assert files["display"][0]["name"] == "ds.txt"
    other = client.post("/items", json={"template_id": templates["FPG"], "values": {
        "description": "reusing someone else's file", "datasheet": [staged.json()["id"]]}},
        headers=admin)
    assert other.status_code == 400


# ─────────────────────────── import / export ───────────────────────────
def _import_file(client, admin, template_id):
    query = f"?template_id={template_id}" if template_id else ""
    r = client.get(f"/data/template{query}", headers=admin)
    assert r.status_code == 200
    return load_workbook(io.BytesIO(r.content))


def _upload(client, headers, wb, name="items.xlsx"):
    buf = io.BytesIO()
    wb.save(buf)
    return client.post("/data/import", files={"file": (name, buf.getvalue(),
                       "application/octet-stream")}, headers=headers)


def test_import_file_headers_are_the_templates_creation_fields(client, admin, templates):
    wb = _import_file(client, admin, templates["FPG"])
    ws = next(s for s in wb.worksheets if s.title.startswith("C-FPG"))
    assert [c.value for c in ws[1]] == ["Serial", "Description", "Board ID", "Letter"]
    wb = _import_file(client, admin, templates["SPM"])
    ws = next(s for s in wb.worksheets if s.title.startswith("A-SPM"))
    assert [c.value for c in ws[1]] == ["Serial", "Rev", "Location", "Contents"]


def test_import_creates_items_and_links_contents(client, admin, templates):
    wb = _import_file(client, admin, None)
    fpga = next(s for s in wb.worksheets if s.title.startswith("C-FPG"))
    fpga.append(["C-FPG-201", "imported fpga one", "00001", "A"])
    fpga.append([None, "imported fpga two", "FP-00002", None])
    spm = next(s for s in wb.worksheets if s.title.startswith("A-SPM"))
    spm.append([None, "Z", "Lab A — Bench 1", "C-FPG-201"])
    r = _upload(client, admin, wb)
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 3
    card = next(i for i in client.get(f"/items?template_id={templates['FPG']}",
                                      headers=admin).json() if i["serial"] == "C-FPG-201")
    assert card["parent_label"].startswith("Signal Processing Module")


def test_import_is_all_or_nothing_and_names_the_bad_cells(client, admin, templates):
    before = len(client.get("/items", headers=admin).json())
    wb = _import_file(client, admin, None)
    fpga = next(s for s in wb.worksheets if s.title.startswith("C-FPG"))
    fpga.append([None, "a perfectly good row", "00003", "B"])
    fpga.append([None, "short", "12", "BB"])          # row 3: three bad cells
    fpga.append(["C-PRB-999", "short", None, None])  # row 4: bad serial *and* description
    r = _upload(client, admin, wb)
    assert r.status_code == 400
    cells = {(e["sheet"][:5], e["cell"]) for e in r.json()["errors"]}
    assert {("C-FPG", "B3"), ("C-FPG", "C3"), ("C-FPG", "D3"), ("C-FPG", "A4"),
            ("C-FPG", "B4")} <= cells, "every bad cell of a row is reported, not just the first"
    assert len(client.get("/items", headers=admin).json()) == before, "nothing was saved"


def test_import_rejects_unknown_columns_and_missing_required_ones(client, admin, templates):
    wb = _import_file(client, admin, templates["FPG"])
    ws = next(s for s in wb.worksheets if s.title.startswith("C-FPG"))
    ws.delete_cols(2)  # drop the required Description column
    ws.cell(row=1, column=5, value="Colour")
    ws.append([None, "00004", "C"])
    r = _upload(client, admin, wb)
    assert r.status_code == 400
    messages = " ".join(e["error"] for e in r.json()["errors"])
    assert "Description" in messages and "Colour" in messages


def test_import_accepts_only_excel_and_only_managers(client, admin, editor):
    r = client.post("/data/import", files={"file": ("x.csv", b"a,b", "text/csv")}, headers=admin)
    assert r.status_code == 400
    r = client.post("/data/import", files={"file": ("x.xlsx", b"not a zip", "x")}, headers=admin)
    assert r.status_code == 400 and "readable" in r.json()["detail"]
    wb = Workbook()
    assert _upload(client, editor, wb).status_code == 403


def test_export_has_a_sheet_per_template(client, admin):
    r = client.get("/data/export", headers=admin)
    wb = load_workbook(io.BytesIO(r.content))
    titles = [ws.title for ws in wb.worksheets]
    assert any(t.startswith("C-PRB") for t in titles)
    prb = next(ws for ws in wb.worksheets if ws.title.startswith("C-PRB"))
    header = [c.value for c in prb[1]]
    assert header[:4] == ["Serial", "State", "Location", "Parent"]
    assert "Project" in header
    first = [c.value for c in prb[2]]
    assert first[header.index("Project")] == "Falcon"


# ─────────────────────────── users, search, map ───────────────────────────
def test_password_over_the_bcrypt_byte_limit_is_a_clean_422(client, admin):
    r = client.post("/users", json={"email": "long@lattice.io", "full_name": "Long",
                                    "password": "א" * 40, "role": "viewer"}, headers=admin)
    assert r.status_code == 422


def test_deleting_a_user_with_proposals_is_refused_not_crashed(client, admin):
    dana = _users(client, admin)["dana@lattice.io"]
    r = client.delete(f"/users/{dana['id']}", headers=admin)
    assert r.status_code == 409


def test_deleting_a_user_cleans_template_references(client, admin):
    r = client.post("/users", json={"email": "temp-mgr@lattice.io", "full_name": "Temp Mgr",
                                    "password": "password", "role": "manager"}, headers=admin)
    uid = r.json()["id"]
    tpl = client.post("/templates", json={
        "type": "setup", "name": "Temp Owned", "serial_prefix": "TMO",
        "fields": [{"label": "Managers", "field_type": "managers", "mode": "fixed",
                    "fixed_value": [uid]}]}, headers=admin).json()
    assert client.delete(f"/users/{uid}", headers=admin).status_code == 204
    t = client.get(f"/templates/{tpl['id']}", headers=admin).json()
    assert t["fields"][0]["fixed_value"] == []


def test_login_hints_expose_only_what_a_manager_published(client):
    hints = client.get("/auth/login-hints").json()
    assert {h["email"] for h in hints} >= {"noa@lattice.io"}
    assert all("admin@lattice.io" != h["email"] for h in hints)


def test_search_finds_serials_and_templates_and_treats_wildcards_literally(client, admin):
    r = client.get("/search?q=C-PRB-00", headers=admin).json()
    assert r["items"] and all("C-PRB-00" in h["title"] for h in r["items"])
    r = client.get("/search?q=regulator", headers=admin).json()
    assert any(h["kind"] == "template" for h in r["templates"])
    assert client.get("/search?q=%25", headers=admin).json()["total"] == 0


def test_map_buildings_crud_and_permissions(client, admin, editor, viewer):
    r = client.post("/map/buildings", json={"name": "Annex", "x": 90, "y": 90, "width": 9,
                                            "height": 9}, headers=editor)
    assert r.status_code == 201
    bid = r.json()["id"]
    assert client.delete(f"/map/buildings/{bid}", headers=editor).status_code == 403
    assert client.delete(f"/map/buildings/{bid}", headers=admin).status_code == 204
    assert client.post("/map/buildings", json={"name": "X"}, headers=viewer).status_code == 403
